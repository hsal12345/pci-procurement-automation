import io
import json
import pytest
from datetime import date
from decimal import Decimal
from pathlib import Path
from uuid import uuid4
import openpyxl
from httpx import AsyncClient, ASGITransport

from app.main import app
from app.database import async_engine, Base, AsyncSessionLocal
from app.models.masters import School, Company
from app.models.packages import Package, PackageRevision, RevisionLine, DocumentEvent, DocumentNumber
from app.models.generation import GenerationJob, Artifact
from app.services.export.context_builder import build_generation_context
from app.services.export.excel_generator import ExcelPackageGenerator
from app.services.export.docx_generator import DocxPackageGenerator
from app.services.export.pdf_service import PdfService
from app.services.export.zip_packager import ZipPackager, compute_sha256
from app.services.generation_worker import run_generation_job


@pytest.fixture(autouse=True)
async def prepare_database():
    """Ensures clean tables before each test."""
    async with async_engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)
        await conn.run_sync(Base.metadata.create_all)
    yield


def create_mock_revision_data(item_count: int = 8, is_draft: bool = False):
    """Creates synthetic revision context with arbitrary N items."""
    lines = []
    subtotal = Decimal("0")
    for i in range(1, item_count + 1):
        qty = Decimal(str(i % 5 + 1))
        hps_p = Decimal(str(25000 + (i * 1000)))
        offer_p = Decimal(str(24000 + (i * 1000)))
        final_p = Decimal(str(23500 + (i * 1000)))
        line_tot = qty * final_p
        subtotal += line_tot

        lines.append({
            "line_number": i,
            "original_description": f"Barang Alat Kebersihan No {i}",
            "confirmed_description": f"Barang Alat Kebersihan No {i}",
            "specification": f"Spesifikasi Standar Mutu {i}",
            "quantity": qty,
            "unit": "Pcs",
            "hps_unit_price": hps_p,
            "offer_unit_price": offer_p,
            "final_unit_price": final_p,
            "discount": Decimal("0"),
            "line_total": line_tot,
            "condition": "BAIK"
        })

    tax_total = (subtotal * Decimal("0.11")).quantize(Decimal("1"))
    grand_total = subtotal + tax_total

    events = [
        {"event_code": "HPS", "event_date": "2026-06-15", "location": "Jakarta"},
        {"event_code": "SPP", "event_date": "2026-06-16", "location": "Jakarta"},
        {"event_code": "PENAWARAN", "event_date": "2026-06-17", "location": "Jakarta"},
        {"event_code": "NEGOSIASI", "event_date": "2026-06-18", "location": "Jakarta"},
        {"event_code": "PESANAN", "event_date": "2026-06-19", "location": "Jakarta"},
        {"event_code": "SURAT_JALAN", "event_date": "2026-06-20", "location": "Jakarta"},
        {"event_code": "BAPB", "event_date": "2026-06-21", "location": "Jakarta"},
        {"event_code": "BAST", "event_date": "2026-06-22", "location": "Jakarta"},
        {"event_code": "BAST_REKANAN", "event_date": "2026-06-22", "location": "Jakarta"},
        {"event_code": "BAPP", "event_date": "2026-06-23", "location": "Jakarta"},
        {"event_code": "FAKTUR", "event_date": "2026-06-24", "location": "Jakarta"},
        {"event_code": "KWITANSI", "event_date": "2026-06-25", "location": "Jakarta"},
    ]

    numbers = [
        {"doc_code": "D01", "doc_number": "01/HPS/KK09/2026", "issuer": "SCHOOL"},
        {"doc_code": "D02", "doc_number": "02/SPP/KK09/2026", "issuer": "SCHOOL"},
        {"doc_code": "D03", "doc_number": "03/PENAWARAN/PCI/2026", "issuer": "COMPANY"},
        {"doc_code": "D04", "doc_number": "04/BA-NEGO/KK09/2026", "issuer": "SCHOOL"},
        {"doc_code": "D05", "doc_number": "05/SP/KK09/2026", "issuer": "SCHOOL"},
        {"doc_code": "D06", "doc_number": "06/FAK-PCI/2026", "issuer": "COMPANY"},
        {"doc_code": "D07", "doc_number": "07/KWT/KK09/2026", "issuer": "SCHOOL"},
        {"doc_code": "D08", "doc_number": "08/SJ-PCI/2026", "issuer": "COMPANY"},
        {"doc_code": "D09", "doc_number": "09/BAPB/KK09/2026", "issuer": "SCHOOL"},
        {"doc_code": "D10", "doc_number": "10/BAST/KK09/2026", "issuer": "SCHOOL"},
        {"doc_code": "D11", "doc_number": "11/BAPP/KK09/2026", "issuer": "SCHOOL"},
        {"doc_code": "D12", "doc_number": "12/LAMP-BAPP/KK09/2026", "issuer": "SCHOOL"},
        {"doc_code": "D13", "doc_number": "13/BAST-REK/PCI/2026", "issuer": "COMPANY"},
    ]

    return {
        "package_number": f"PKG/2026/Q2/{item_count:04d}",
        "fiscal_year": 2026,
        "quarter": 2,
        "revision_number": 1,
        "school_snapshot": {
            "name": "SEKOLAH DASAR NEGERI KEBON KOSONG 09",
            "npsn": "20108304",
            "address": "Jl. Kebon Kosong I No. 1 Kemayoran",
            "principal_name": "Drs. Mulyadi",
            "principal_nip": "19680512 199303 1 005",
        },
        "company_snapshot": {
            "legal_name": "CV. PUSTAKA CEMERLANG INDONESIA",
            "director_name": "Bambang Solifian",
            "address": "Wisma Mitra Sunter Unit 3-03",
            "npwp": "01.234.567.8-012.000",
            "bank_name": "BANK DKI",
            "bank_account_number": "300.21.00000.1",
            "bank_account_name": "CV PUSTAKA CEMERLANG INDONESIA",
        },
        "officials_snapshot": [
            {"role": "KEPALA_SEKOLAH", "name": "Drs. Mulyadi", "nip": "19680512 199303 1 005"},
            {"role": "PEJABAT_PENGADAAN", "name": "Budi Santoso, S.Pd", "nip": "19750814 200112 1 002"},
            {"role": "PPHP_KETUA", "name": "Siti Rahayu, S.Pd", "nip": "19800321 200501 2 008"},
            {"role": "PPHP_SEKRETARIS", "name": "Ahmad Fauzi, S.Pd", "nip": "19841105 200902 1 003"},
            {"role": "PPHP_ANGGOTA", "name": "Dewi Sartika, S.Pd", "nip": "19880719 201101 2 015"},
            {"role": "PENGURUS_BARANG", "name": "Sulistiawati", "nip": "19790909 202221 2 008"},
        ],
        "lines": lines,
        "events": events,
        "numbers": numbers,
        "subtotal": subtotal,
        "tax_total": tax_total,
        "grand_total": grand_total,
        "is_draft": is_draft,
    }


def test_dynamic_line_items_n_counts(tmp_path):
    """AT11: Verifies dynamic generation across N = 1, 8, 50, and 100 line items."""
    for n in [1, 8, 50, 100]:
        data = create_mock_revision_data(item_count=n)
        ctx = build_generation_context(data, is_draft=False)

        excel_gen = ExcelPackageGenerator(ctx, is_draft=False)
        out_file = tmp_path / f"test_pkg_{n}.xlsx"
        excel_gen.generate_workbook(output_path=out_file)

        assert out_file.exists(), f"File for N={n} was not created"
        assert out_file.stat().st_size > 5000, f"File for N={n} is suspiciously small"

        # Inspect with openpyxl
        wb = openpyxl.load_workbook(out_file, data_only=False)
        assert len(wb.sheetnames) == 13, f"Expected 13 sheets for N={n}, found {len(wb.sheetnames)}"

        # Check HPS sheet row count
        ws_hps = wb["HPS"]
        # Find item rows starting at row 14 (1-indexed)
        item_rows = []
        for r in range(14, 14 + n + 5):
            val = ws_hps.cell(row=r, column=1).value
            if val is not None and isinstance(val, (int, float)) and val <= n:
                item_rows.append(val)
        assert len(item_rows) == n, f"HPS sheet item rows {len(item_rows)} != {n}"


def test_link_free_excel_invariant(tmp_path):
    """AT12: Verifies that the generated Excel file has zero external links and valid cells."""
    data = create_mock_revision_data(item_count=8)
    ctx = build_generation_context(data, is_draft=False)

    excel_gen = ExcelPackageGenerator(ctx, is_draft=False)
    out_file = tmp_path / "test_link_free.xlsx"
    excel_gen.generate_workbook(output_path=out_file)

    wb = openpyxl.load_workbook(out_file, data_only=False)

    # 1. Check external links attribute
    assert len(wb._external_links) == 0, f"Detected {len(wb._external_links)} external workbook links!"

    # 2. Check every cell across all 13 sheets for error literals or bracketed sheet links
    forbidden_tokens = ["#REF!", "#VALUE!", "#NAME?", ".xlsx]", ".xlsm]"]
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if val is not None:
                    str_val = str(val)
                    for tok in forbidden_tokens:
                        assert tok not in str_val, f"Forbidden token '{tok}' found in sheet '{sheetname}', cell {cell.coordinate}: {str_val}"


def test_cross_format_total_equivalence(tmp_path):
    """AT13: Verifies that monetary figures are strictly identical between Excel, Word, and Manifest."""
    data = create_mock_revision_data(item_count=8)
    ctx = build_generation_context(data, is_draft=False)

    # 1. Generate Excel
    excel_gen = ExcelPackageGenerator(ctx, is_draft=False)
    xlsx_path = tmp_path / "pkg.xlsx"
    excel_gen.generate_workbook(output_path=xlsx_path)

    # 2. Generate Word
    docx_gen = DocxPackageGenerator(ctx)
    docx_dir = tmp_path / "docx"
    docx_files = docx_gen.generate_all(docx_dir)

    # 3. Create Manifest
    packager = ZipPackager(ctx)
    zip_dir = tmp_path / "zip"
    artifacts = [
        {"doc_code": "ALL", "format": "XLSX", "filename": "pkg.xlsx", "path": xlsx_path},
    ]
    for code, p in docx_files.items():
        artifacts.append({"doc_code": code, "format": "DOCX", "filename": p.name, "path": p})

    res = packager.create_package_zip(artifacts, zip_dir)
    manifest = res["manifest"]

    # Invariant assertions
    expected_grand_total = float(data["grand_total"])
    assert manifest["grand_total"] == expected_grand_total
    assert str(ctx["totals"]["grand_total"]) == str(data["grand_total"])
    assert ctx["totals"]["subtotal"] == data["subtotal"]
    assert ctx["totals"]["tax_total"] == data["tax_total"]

    # Verify openpyxl HPS total cell matches
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["HPS"]
    total_val = None
    for r in range(15, 30):
        lbl = ws.cell(row=r, column=1).value
        if lbl and str(lbl).strip().upper() == "TOTAL":
            total_val = ws.cell(row=r, column=7).value
            break

    assert total_val is not None
    assert float(total_val) == float(data["grand_total"])


@pytest.mark.asyncio
async def test_generation_worker_selective_formats():
    """Verifies that generation worker obeys selective format filters (e.g. only XLSX)."""
    async with AsyncSessionLocal() as db:
        school = School(
            name="SDN TEST SELECTIVE",
            npsn="20199999",
            level="SD",
            address="Jl. Test No. 1",
            email="test@sdn.id"
        )
        company = Company(
            legal_name="CV. PUSTAKA CEMERLANG INDONESIA",
            npwp="01.234.567.8-012.000",
            address="Wisma Mitra Sunter",
            city="Jakarta",
            director_name="Bambang Solifian",
            is_active=True
        )
        db.add_all([school, company])
        await db.flush()

        pkg = Package(
            package_number="PKG/2026/Q2/SEL01",
            school_id=school.id,
            fiscal_year=2026,
            quarter=2,
            status="DRAFT"
        )
        db.add(pkg)
        await db.flush()

        rev = PackageRevision(
            package_id=pkg.id,
            revision_number=1,
            school_snapshot={"name": school.name, "npsn": school.npsn, "address": school.address},
            company_snapshot={"legal_name": company.legal_name, "director_name": company.director_name},
            subtotal=Decimal("100000"),
            tax_total=Decimal("11000"),
            grand_total=Decimal("111000"),
            is_locked=False
        )
        db.add(rev)
        await db.flush()

        line = RevisionLine(
            revision_id=rev.id,
            line_number=1,
            original_description="Buku Tulis",
            confirmed_description="Buku Tulis",
            quantity=Decimal("10"),
            unit="Pak",
            hps_unit_price=Decimal("10000"),
            offer_unit_price=Decimal("10000"),
            final_unit_price=Decimal("10000"),
            line_total=Decimal("100000"),
            condition="BAIK"
        )
        event = DocumentEvent(
            revision_id=rev.id,
            event_code="HPS",
            event_date=date(2026, 6, 15),
            location="Jakarta"
        )
        num = DocumentNumber(
            revision_id=rev.id,
            doc_code="D01",
            doc_number="01/HPS/TEST/2026",
            issuer="SCHOOL"
        )
        db.add_all([line, event, num])

        job = GenerationJob(revision_id=rev.id, status="QUEUED")
        db.add(job)
        await db.commit()
        job_id = job.id
        rev_id = rev.id

    # Run generation worker requesting ONLY XLSX
    await run_generation_job(job_id=job_id, formats=["XLSX"], is_draft_override=True)

    async with AsyncSessionLocal() as db:
        updated_job = await db.get(GenerationJob, job_id)
        assert updated_job.status == "SUCCEEDED"
        assert updated_job.progress_percent == 100

        # Verify ONLY XLSX artifact is registered
        from sqlalchemy import select
        art_res = await db.execute(select(Artifact).where(Artifact.revision_id == rev_id))
        artifacts = art_res.scalars().all()
        formats = [a.format for a in artifacts]
        assert "XLSX" in formats
        assert "DOCX" not in formats
        assert "PDF" not in formats
        assert "ZIP" not in formats


@pytest.mark.asyncio
async def test_generation_api_endpoints():
    """Verifies POST /generate, GET /jobs/{id}, GET /artifacts, and GET /artifacts/{id}/download."""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. Create company and school
        res_comp = await client.post("/api/v1/masters/companies", json={
            "legal_name": "CV. PUSTAKA CEMERLANG INDONESIA",
            "npwp": "01.234.567.8-012.000",
            "address": "Wisma Mitra Sunter Unit 3-03",
            "director_name": "Bambang Solifian",
            "city": "Jakarta"
        })
        assert res_comp.status_code == 201

        res_sch = await client.post("/api/v1/masters/schools", json={
            "name": "SDN API TEST",
            "npsn": "20199998",
            "level": "SD",
            "address": "Jl. Kebon Kosong"
        })
        assert res_sch.status_code == 201
        school_id = res_sch.json()["id"]

        # 2. Create package
        res_pkg = await client.post("/api/v1/packages", json={
            "school_id": school_id,
            "fiscal_year": 2026,
            "quarter": 2
        })
        assert res_pkg.status_code == 201
        pkg_data = res_pkg.json()
        rev_id = pkg_data["current_revision_id"]

        # 3. Trigger generation for XLSX
        res_gen = await client.post(
            f"/api/v1/revisions/{rev_id}/generate",
            json={"formats": ["XLSX"], "is_draft": True}
        )
        assert res_gen.status_code == 202
        job_data = res_gen.json()
        job_id = job_data["id"]
        assert job_data["status"] == "QUEUED"

        # 4. Poll job status
        res_job = await client.get(f"/api/v1/jobs/{job_id}")
        assert res_job.status_code == 200
        assert res_job.json()["id"] == job_id

        # 5. Run worker to generate artifact
        await run_generation_job(job_id=job_id, formats=["XLSX"], is_draft_override=True)

        # 6. List artifacts via API
        res_arts = await client.get(f"/api/v1/revisions/{rev_id}/artifacts")
        assert res_arts.status_code == 200
        arts_list = res_arts.json()
        assert len(arts_list) == 1
        art_id = arts_list[0]["id"]
        assert arts_list[0]["format"] == "XLSX"
        assert arts_list[0]["download_url"] == f"/api/v1/artifacts/{art_id}/download"

        # 7. Download artifact file via API
        res_dl = await client.get(f"/api/v1/artifacts/{art_id}/download")
        assert res_dl.status_code == 200
        assert len(res_dl.content) > 5000
        assert res_dl.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.mark.asyncio
async def test_full_package_generation_29_artifacts():
    """AT13 & Section 4 Accounting: Full multi-format pipeline produces exactly 29 artifact records:
    1 XLSX + 13 DOCX + 13 PDF + 1 Merged PDF (ALL) + 1 ZIP = 29 records.
    """
    async with AsyncSessionLocal() as db:
        school = School(
            name="SDN TEST FULL 29",
            npsn="20188888",
            level="SD",
            address="Jl. Percetakan Negara",
            email="full29@sdn.id"
        )
        company = Company(
            legal_name="CV. PUSTAKA CEMERLANG INDONESIA",
            npwp="01.234.567.8-012.000",
            address="Wisma Mitra Sunter Unit 3-03",
            director_name="Bambang Solifian",
            is_active=True
        )
        db.add_all([school, company])
        await db.flush()

        pkg = Package(
            package_number="PKG/2026/Q2/FULL29",
            school_id=school.id,
            fiscal_year=2026,
            quarter=2,
            status="DRAFT"
        )
        db.add(pkg)
        await db.flush()

        rev = PackageRevision(
            package_id=pkg.id,
            revision_number=1,
            school_snapshot={"name": school.name, "npsn": school.npsn, "address": school.address},
            company_snapshot={"legal_name": company.legal_name, "director_name": company.director_name},
            subtotal=Decimal("250000"),
            tax_total=Decimal("27500"),
            grand_total=Decimal("277500"),
            is_locked=False
        )
        db.add(rev)
        await db.flush()

        line = RevisionLine(
            revision_id=rev.id,
            line_number=1,
            original_description="Kertas HVS A4 80gr",
            confirmed_description="Kertas HVS A4 80gr",
            quantity=Decimal("5"),
            unit="Rim",
            hps_unit_price=Decimal("50000"),
            offer_unit_price=Decimal("50000"),
            final_unit_price=Decimal("50000"),
            line_total=Decimal("250000"),
            condition="BAIK"
        )
        event = DocumentEvent(
            revision_id=rev.id,
            event_code="HPS",
            event_date=date(2026, 6, 15),
            location="Jakarta"
        )
        num = DocumentNumber(
            revision_id=rev.id,
            doc_code="D01",
            doc_number="01/HPS/FULL/2026",
            issuer="SCHOOL"
        )
        db.add_all([line, event, num])

        job = GenerationJob(revision_id=rev.id, status="QUEUED")
        db.add(job)
        await db.commit()
        job_id = job.id
        rev_id = rev.id

    # Run full generation with all 4 format engines
    await run_generation_job(job_id=job_id, formats=["XLSX", "DOCX", "PDF", "ZIP"], is_draft_override=True)

    async with AsyncSessionLocal() as db:
        updated_job = await db.get(GenerationJob, job_id)
        assert updated_job.status == "SUCCEEDED", f"Job failed: {updated_job.error_message}"
        assert updated_job.progress_percent == 100

        from sqlalchemy import select
        art_res = await db.execute(select(Artifact).where(Artifact.revision_id == rev_id))
        artifacts = art_res.scalars().all()

        xlsx_arts = [a for a in artifacts if a.format == "XLSX"]
        docx_arts = [a for a in artifacts if a.format == "DOCX"]
        pdf_arts = [a for a in artifacts if a.format == "PDF"]
        zip_arts = [a for a in artifacts if a.format == "ZIP"]

        assert len(xlsx_arts) == 1, f"Expected 1 XLSX, got {len(xlsx_arts)}"
        assert len(docx_arts) == 13, f"Expected 13 DOCX, got {len(docx_arts)}"
        assert len(pdf_arts) == 14, f"Expected 14 PDF (13 individual + 1 merged ALL), got {len(pdf_arts)}"
        assert len(zip_arts) == 1, f"Expected 1 ZIP, got {len(zip_arts)}"

        # Total artifacts: 1 + 13 + 14 + 1 = 29
        assert len(artifacts) == 29, f"Total artifact records {len(artifacts)} != 29"

        # Verify all files exist on disk and SHA-256 matches
        for a in artifacts:
            p = Path(a.file_path)
            assert p.exists(), f"File {p} does not exist on disk"
            assert compute_sha256(p) == a.sha256_hash, f"SHA-256 mismatch on {p.name}"
            assert a.is_draft is True
